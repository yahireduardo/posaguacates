#!/usr/bin/env python3
"""Extrae cajas de todas las clasificaciones GRANDES de libros históricos."""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


MONTHS = {
    "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4,
    "MAYO": 5, "JUNIO": 6, "JULIO": 7, "AGOSTO": 8,
    "SEPTIEMBRE": 9, "SETIEMBRE": 9, "OCTUBRE": 10,
    "NOVIEMBRE": 11, "DICIEMBRE": 12,
}


def normalize(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").upper())
    return " ".join("".join(char for char in text if not unicodedata.combining(char)).split())


def parse_date(*texts: object, default_year: int | None = None) -> date | None:
    combined = " ".join(normalize(text) for text in texts if text)
    match = re.search(r"\b(\d{1,2})\s+(?:DE\s+)?(" + "|".join(MONTHS) + r")(?:\s+(\d{4}))?\b", combined)
    if not match:
        return None
    year_match = match.group(3) or re.search(r"\b(20\d{2})\b", combined)
    if hasattr(year_match, "group"):
        year_value = year_match.group(1)
    else:
        year_value = year_match or default_year
    if not year_value:
        return None
    year = int(year_value)
    try:
        return date(year, MONTHS[match.group(2)], int(match.group(1)))
    except ValueError:
        return None


def is_large_classification(value: object) -> bool:
    text = normalize(value)
    return bool(re.search(r"\bGRANDES?\b", text))


def numeric(value: object) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def extract_file(path: Path) -> tuple[list[dict], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    records, warnings = [], []
    filename_year = re.search(r"\b(20\d{2})\b", path.name)
    default_year = int(filename_year.group(1)) if filename_year else None
    try:
        for worksheet in workbook.worksheets:
            header = worksheet.cell(1, 1).value
            columns = [normalize(worksheet.cell(3, column).value) for column in range(1, min(worksheet.max_column, 12) + 1)]
            if not header or "AGUACATE HASS" not in normalize(header) or "CLASIFICACION" not in columns:
                continue
            batch_date = parse_date(worksheet.cell(2, 1).value, header, worksheet.title, default_year=default_year)
            if not batch_date:
                warnings.append(f"{path.name} / {worksheet.title}: no se reconoció la fecha")
                continue
            if batch_date > date.today():
                warnings.append(f"{path.name} / {worksheet.title}: fecha futura excluida ({batch_date.isoformat()})")
                continue
            classifications, total = [], 0.0
            for row in range(4, min(worksheet.max_row, 100) + 1):
                label = worksheet.cell(row, 1).value
                if not is_large_classification(label):
                    continue
                boxes = numeric(worksheet.cell(row, 2).value)
                if boxes is None or boxes <= 0:
                    continue
                total += boxes
                classifications.append(f"{normalize(label)}: {boxes:g}")
            if total > 0:
                records.append({
                    "date": batch_date.isoformat(), "file": path.name,
                    "sheet": worksheet.title, "quantity": round(total, 2),
                    "classifications": classifications,
                })
    finally:
        workbook.close()
    return records, warnings


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("files", nargs="+")
    args = parser.parse_args()
    all_records, all_warnings = [], []
    for filename in args.files:
        path = Path(filename)
        if not path.is_file():
            all_warnings.append(f"No existe: {path}")
            continue
        records, warnings = extract_file(path)
        all_records.extend(records)
        all_warnings.extend(warnings)
    json.dump({"records": all_records, "warnings": all_warnings}, sys.stdout, ensure_ascii=False)


if __name__ == "__main__":
    main()
